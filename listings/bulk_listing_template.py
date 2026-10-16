"""Build the small CSV/XLSX templates used by the bulk-listing importer.

The importer intentionally accepts only one fixed header and at most twenty
rows.  Rather than maintaining another static asset, this module generates that
contract on demand from the parser constants.  Nothing here reads seller data,
touches the database, or stores a file.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .bulk_listing_csv import (
    BULK_LISTING_IMPORT_HEADERS,
    BULK_LISTING_IMPORT_MAX_ROWS,
    TRADE_MODE_LABELS,
)

TEMPLATE_FORMATS = ('csv', 'xlsx')


@dataclass(frozen=True)
class BulkListingImportTemplate:
    """A complete downloadable response payload."""

    content: bytes
    content_type: str
    filename: str


def _build_csv_template() -> BulkListingImportTemplate:
    buffer = io.StringIO(newline='')
    writer = csv.writer(buffer, lineterminator='\r\n')
    writer.writerow(BULK_LISTING_IMPORT_HEADERS)
    return BulkListingImportTemplate(
        content=buffer.getvalue().encode('utf-8'),
        content_type='text/csv; charset=utf-8',
        filename='campus-glimmer-bulk-listing-template.csv',
    )


def _build_xlsx_template() -> BulkListingImportTemplate:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '商品导入'
    worksheet.append(BULK_LISTING_IMPORT_HEADERS)

    worksheet.freeze_panes = 'A2'
    for column, width in zip('ABCDEFG', (30, 46, 16, 14, 14, 16, 18)):
        worksheet.column_dimensions[column].width = width
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color='FFFFFFFF')
        cell.fill = PatternFill('solid', fgColor='2C5F82')
        cell.alignment = Alignment(horizontal='center')

    trade_choices = ','.join(TRADE_MODE_LABELS)
    validation = DataValidation(
        type='list',
        formula1=f'"{trade_choices}"',
        allow_blank=True,
        showErrorMessage=True,
        errorTitle='交易方式',
        error=f'只能填写：{trade_choices}',
    )
    last_row = BULK_LISTING_IMPORT_MAX_ROWS + 1
    validation.add(f'C2:C{last_row}')
    worksheet.add_data_validation(validation)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return BulkListingImportTemplate(
        content=buffer.getvalue(),
        content_type=(
            'application/vnd.openxmlformats-officedocument.'
            'spreadsheetml.sheet'
        ),
        filename='campus-glimmer-bulk-listing-template.xlsx',
    )


_TEMPLATE_BUILDERS = {
    'csv': _build_csv_template,
    'xlsx': _build_xlsx_template,
}


def build_bulk_listing_import_template(template_format: str) -> BulkListingImportTemplate:
    """Return one allowed template; callers turn anything else into a 404."""
    try:
        builder = _TEMPLATE_BUILDERS[template_format]
    except KeyError as exc:
        raise ValueError(f'unsupported template format: {template_format}') from exc
    return builder()