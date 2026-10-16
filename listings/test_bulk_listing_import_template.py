from io import BytesIO
from zipfile import ZipFile

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .bulk_listing_csv import BULK_LISTING_IMPORT_HEADERS, TRADE_MODE_LABELS, parse_bulk_listing_csv
from .bulk_listing_template import build_bulk_listing_import_template
from .bulk_listing_xlsx import parse_bulk_listing_xlsx


class BulkListingImportTemplateTests(TestCase):
    """批量发布 CSV / XLSX 模板的生成与下载边界。"""

    @classmethod
    def setUpTestData(cls):
        cls.seller = User.objects.create_user(
            username='template-seller', password='safe-password-123',
        )

    def test_csv_download_uses_exact_header_and_safe_headers(self):
        self.client.force_login(self.seller)
        page = self.client.get(reverse('bulk_create_items'))
        self.assertEqual(page.status_code, 200)
        self.assertContains(
            page,
            reverse('bulk_listing_import_template', args=['csv']),
        )

        response = self.client.get(
            reverse('bulk_listing_import_template', args=['csv']),
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'text/csv; charset=utf-8')
        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="campus-glimmer-bulk-listing-template.csv"',
        )
        self.assertEqual(response['Cache-Control'], 'no-store')
        self.assertFalse(response.content.startswith(b'\xef\xbb\xbf'))
        self.assertEqual(
            response.content.decode('utf-8').splitlines(),
            [','.join(BULK_LISTING_IMPORT_HEADERS)],
        )

    def test_xlsx_download_contains_one_guarded_worksheet(self):
        self.client.force_login(self.seller)

        response = self.client.get(
            reverse('bulk_listing_import_template', args=['xlsx']),
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            'spreadsheetml.sheet', response['Content-Type'],
        )
        self.assertEqual(
            response['Content-Disposition'],
            'attachment; filename="campus-glimmer-bulk-listing-template.xlsx"',
        )
        self.assertEqual(response['Cache-Control'], 'no-store')

        with ZipFile(BytesIO(response.content)) as archive:
            self.assertIsNone(archive.testzip())
        workbook = load_workbook(BytesIO(response.content), data_only=False)
        self.assertEqual(workbook.sheetnames, ['商品导入'])
        worksheet = workbook.active
        self.assertEqual(
            tuple(cell.value for cell in worksheet[1]),
            BULK_LISTING_IMPORT_HEADERS,
        )
        self.assertEqual(worksheet.freeze_panes, 'A2')
        self.assertEqual(len(worksheet.data_validations.dataValidation), 1)
        validation = worksheet.data_validations.dataValidation[0]
        self.assertEqual(str(validation.sqref), 'C2:C21')
        self.assertEqual(
            validation.formula1, f'"{",".join(TRADE_MODE_LABELS)}"',
        )
        self.assertTrue(all(cell.data_type == 's' for cell in worksheet[1]))

    def test_generated_templates_still_reach_the_empty_row_guard(self):
        csv_template = build_bulk_listing_import_template('csv')
        csv_result = parse_bulk_listing_csv(SimpleUploadedFile(
            csv_template.filename,
            csv_template.content,
            content_type='text/csv',
        ))
        self.assertEqual(csv_result.file_error, '')
        self.assertEqual(len(csv_result.row_errors), 1)
        self.assertIn('CSV 中至少需要一行非空商品', csv_result.row_errors[0].message)
        self.assertEqual(csv_result.rows, [])

        xlsx_template = build_bulk_listing_import_template('xlsx')
        xlsx_result = parse_bulk_listing_xlsx(SimpleUploadedFile(
            xlsx_template.filename,
            xlsx_template.content,
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'
            ),
        ))
        self.assertEqual(xlsx_result.file_error, '')
        self.assertEqual(len(xlsx_result.row_errors), 1)
        self.assertIn('至少需要一行非空商品', xlsx_result.row_errors[0].message)
        self.assertEqual(xlsx_result.rows, [])

    def test_unknown_format_and_anonymous_download_are_rejected(self):
        self.client.force_login(self.seller)
        unknown = self.client.get(
            reverse('bulk_listing_import_template', args=['xls']),
        )
        self.assertEqual(unknown.status_code, 404)
        self.client.logout()

        response = self.client.get(
            reverse('bulk_listing_import_template', args=['csv']),
        )
        self.assertEqual(response.status_code, 302)
        self.assertIn('login', response['Location'])
